"""Golden set: a blind, hand-labelled sample and agreement scoring against the AI.

  python run.py golden-sheet   -> data/eval/golden_worksheet.xlsx (no AI labels shown)
  (you label it by hand in Excel / Google Sheets, save as golden_labels.xlsx or .csv)
  python run.py golden-score   -> data/eval/golden_report.json + per-item disagreements
"""
import hashlib
import json
import random
from collections import Counter

from engine import codebook as CB
from engine import config as C
from engine.prefilter import load_corpus

GOLD_FIELDS = ["relevance", "product", "investor_experience", "journey_stage",
               "drop_trigger", "unmet_need", "continuation_risk"]
SHEET = C.EVAL / "golden_worksheet.xlsx"


def sample(corpus, n=20, seed=7):
    """Source-stratified random sample. Uses no AI output, so the sample is blind."""
    rng = random.Random(seed)
    by_src = {}
    for r in corpus:
        by_src.setdefault(r["source"], []).append(r)
    for v in by_src.values():
        v.sort(key=lambda r: hashlib.md5(r["id"].encode()).hexdigest())
        rng.shuffle(v)
    picked, i = [], 0
    while len(picked) < n and any(by_src.values()):
        src = sorted(by_src)[i % len(by_src)]
        if by_src[src]:
            r = by_src[src].pop()
            if r["thread_id"] not in {p["thread_id"] for p in picked}:
                picked.append(r)
        i += 1
    return picked


def make_sheet(n=20):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    rows = sample(load_corpus(), n)
    wb = Workbook()
    ws = wb.active
    ws.title = "label_here"
    head = ["id", "source", "url", "context", "title", "text"] + GOLD_FIELDS + ["your_notes"]
    ws.append(head)
    for r in rows:
        ws.append([r["id"], r["source"], r["url"], r.get("context", ""), r.get("title", ""),
                   r["text"][:1500]] + [""] * (len(GOLD_FIELDS) + 1))
    lists = wb.create_sheet("allowed_values")
    for col, f in enumerate(GOLD_FIELDS, 1):
        lists.cell(1, col, f).font = Font(bold=True)
        for i, v in enumerate(CB.FIELDS[f]["values"], 2):
            lists.cell(i, col, v)
        letter = lists.cell(1, col).column_letter
        dv = DataValidation(type="list", formula1=f"=allowed_values!${letter}$2:${letter}${len(CB.FIELDS[f]['values']) + 1}",
                            allow_blank=True)
        ws.add_data_validation(dv)
        tcol = ws.cell(1, 6 + col).column_letter
        dv.add(f"{tcol}2:{tcol}{len(rows) + 1}")
    guide = wb.create_sheet("codebook")
    guide.append(["field", "value", "meaning"])
    for f in GOLD_FIELDS:
        for v, d in CB.FIELDS[f]["values"].items():
            guide.append([f, v, d])
    fill = PatternFill("solid", fgColor="FFF2CC")
    for c in ws[1]:
        c.font = Font(bold=True)
    for col in range(7, 7 + len(GOLD_FIELDS)):
        for row in range(1, len(rows) + 2):
            ws.cell(row, col).fill = fill
    for col, w in zip("ABCDEF", (14, 11, 30, 30, 30, 80)):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(SHEET)
    (C.EVAL / "golden_ids.json").write_text(json.dumps([r["id"] for r in rows], indent=1))
    print(f"golden worksheet with {len(rows)} items -> {SHEET}")
    return SHEET


def normalise(field, v):
    """Forgive human typing: case, stray spaces/newlines, blanks, 2.0 vs 2."""
    import re
    v = "" if v is None or str(v).lower() == "nan" else str(v)
    v = re.sub(r"\s*_\s*", "_", v.strip().lower())
    v = re.sub(r"\s+", "_", v)
    if v.endswith(".0"):
        v = v[:-2]
    if not v:
        vals = CB.FIELDS[field]["values"]
        v = "none" if "none" in vals else ("unknown" if "unknown" in vals else "0")
    return v


def read_labels(path=None):
    import pandas as pd
    for p in ([path] if path else [C.EVAL / "golden_labels.xlsx", C.EVAL / "golden_labels.csv"]):
        if p and p.exists():
            df = pd.read_csv(p) if p.suffix == ".csv" else pd.read_excel(p, sheet_name=0)
            out = {}
            for _, r in df.iterrows():
                if all(str(r.get(f, "")).lower() in ("", "nan") for f in GOLD_FIELDS):
                    continue
                out[str(r["id"])] = {f: normalise(f, r.get(f)) for f in GOLD_FIELDS}
            bad = {(f, v) for h in out.values() for f, v in h.items() if v not in CB.FIELDS[f]["values"]}
            if bad:
                print("  ! labels not in the codebook (counted as written):", sorted(bad))
            return out
    raise FileNotFoundError("Save your labelled sheet as data/eval/golden_labels.xlsx (or .csv)")


def golden_ids():
    """The ids the human actually labelled (falls back to the sheet that was generated)."""
    try:
        return list(read_labels())
    except FileNotFoundError:
        p = C.EVAL / "golden_ids.json"
        return json.loads(p.read_text()) if p.exists() else []


def kappa(pairs):
    if not pairs:
        return None
    n = len(pairs)
    po = sum(a == b for a, b in pairs) / n
    ca, cb = Counter(a for a, _ in pairs), Counter(b for _, b in pairs)
    pe = sum(ca[k] * cb.get(k, 0) for k in ca) / (n * n)
    return round((po - pe) / (1 - pe), 3) if pe < 1 else 1.0


def score(path=None):
    from engine.classify import load_done
    human = read_labels(path)
    ai = load_done()
    corpus = {r["id"]: r for r in load_corpus()}
    # records the code prefilter dropped count as the engine saying "not relevant"
    from engine.prefilter import load_raw
    for r in load_raw():
        if r["id"] in human and r["id"] not in corpus:
            corpus[r["id"]] = r
            ai.setdefault(r["id"], {"relevance": "not_relevant", "product": "general",
                                    "investor_experience": "unknown", "journey_stage": "unknown",
                                    "drop_trigger": "none", "unmet_need": "none",
                                    "continuation_risk": "0", "status": "prefilter_excluded",
                                    "model": "code prefilter", "reason": "no SIP/MF terms"})
    per_field, items = {}, []
    for f in GOLD_FIELDS:
        pairs = []
        for i, h in human.items():
            if not h.get(f):
                continue
            a = ai.get(i)
            av = (a or {}).get(f, "MISSING") if a and a.get("status") != "rejected" else "REJECTED/MISSING"
            pairs.append((h[f], str(av)))
        agree = sum(a == b for a, b in pairs)
        per_field[f] = {"n": len(pairs), "agree": agree,
                        "agreement": round(agree / len(pairs), 3) if pairs else None,
                        "kappa": kappa(pairs)}
    for i, h in human.items():
        a = ai.get(i, {})
        diffs = {f: {"human": h[f], "ai": a.get(f, "MISSING")} for f in GOLD_FIELDS
                 if h.get(f) and str(a.get(f, "MISSING")) != h[f]}
        rec = corpus.get(i, {})
        items.append({"id": i, "source": rec.get("source"), "url": rec.get("url"),
                      "text": rec.get("text", "")[:400], "ai_quote": a.get("quote"),
                      "ai_reason": a.get("reason"), "ai_model": a.get("model"),
                      "ai_status": a.get("status", "not coded"),
                      "all_fields_agree": not diffs, "disagreements": diffs})
    by_model = {}
    for it in items:
        m = by_model.setdefault(it["ai_model"] or "not coded", {"items": 0, "pairs": 0, "agree": 0})
        m["items"] += 1
        h = human[it["id"]]
        for f in GOLD_FIELDS:
            if h.get(f):
                m["pairs"] += 1
                m["agree"] += f not in it["disagreements"]
    for m in by_model.values():
        m["agreement"] = round(m["agree"] / m["pairs"], 3) if m["pairs"] else None
    total_pairs = sum(v["n"] for v in per_field.values())
    total_agree = sum(v["agree"] for v in per_field.values())
    report = {"items_labelled": len(human),
              "items_fully_agreed": sum(x["all_fields_agree"] for x in items),
              "field_level_agreement": round(total_agree / total_pairs, 3) if total_pairs else None,
              "per_field": per_field, "by_model": by_model, "items": items}
    (C.EVAL / "golden_report.json").write_text(json.dumps(report, indent=1, ensure_ascii=False))
    print(f"golden: {report['items_labelled']} items, field agreement "
          f"{report['field_level_agreement']}, fully agreed items {report['items_fully_agreed']}")
    for f, v in per_field.items():
        print(f"  {f:22s} {v['agree']}/{v['n']}  kappa={v['kappa']}")
    for m, v in by_model.items():
        print(f"  by labeller: {m:32s} {v['items']} items, agreement {v['agreement']}")
    return report
