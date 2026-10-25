"""End-to-end pipeline test on SYNTHETIC records with a fake LLM (no network, no key).

    python -m unittest discover -s tests -v
"""
import json
import os
import random
import tempfile
import unittest
from pathlib import Path

TMP = tempfile.mkdtemp()
os.environ["DISCOVERY_DATA"] = TMP

from engine import ask as ASK  # noqa: E402
from engine import classify as CL  # noqa: E402
from engine import codebook as CB  # noqa: E402
from engine import config as C  # noqa: E402

SYN = [
    ("play_store", "My first SIP of 1000 failed because autopay mandate expired and nobody told me, so it lapsed"),
    ("reddit", "Started my first SIP 6 months ago, returns are only 2%. Should I stop? Expected more"),
    ("reddit", "Lost my job, can I pause my SIP for 3 months without breaking it?"),
    ("app_store", "XIRR shown in the app is confusing, total returns look red, don't understand mutual fund numbers"),
    ("youtube", "I kept my SIP going for 3 years because it is automatic and I never check it"),
    ("tradingqna", "How do I skip one SIP instalment? the app only shows cancel option"),
    ("play_store", "Great app for US stocks, smooth UI"),
    ("reddit", "Friend told me to stop my SIP and move to FD after market fell, first time investor here"),
]


class FakeLLM:
    available = True
    calls = tokens = 0
    seconds = 0.0
    exhausted = set()

    def __init__(self, break_quote=False):
        self.break_quote = break_quote

    def chat_json(self, system, user, max_tokens=0):
        recs = json.loads(user[user.index("{"):])["records"]
        items = []
        for r in recs:
            t = r["text"].lower()
            rel = "not_relevant" if "sip" not in t and "mutual" not in t else "core"
            stop = any(w in t for w in ("stop", "lapsed", "pause"))
            item = {f: ([next(iter(v["values"]))] if v["multi"] else "unknown")
                    for f, v in CB.FIELDS.items()}
            item.update(id=r["id"], relevance=rel, product="indmoney",
                        investor_experience="first_time_explicit" if "first" in t else "unknown",
                        journey_stage="considering_stop" if stop else ("ongoing" if rel == "core" else "unknown"),
                        drop_trigger="payment_mandate_failure" if "mandate" in t else ("income_cashflow" if "job" in t else "none"),
                        continuation_risk="2" if stop else "0",
                        unmet_need="reliable_payments" if "mandate" in t else ("flexible_control" if "pause" in t or "skip" in t else "understand_volatility"),
                        in_app_pain=["none"], start_motivation=["unknown"], what_checked=["unknown"],
                        outside_app=["none"], continuation_driver=["none"], check_frequency="unknown",
                        expectation="unknown", income_stability="unknown", age_band="unknown",
                        money_source="unknown", influencer="unknown",
                        need_text="know what happens", reason="synthetic",
                        quote=("NOT IN TEXT" if self.break_quote else r["text"][:20]))
            if rel == "not_relevant":
                item.update(journey_stage="unknown", drop_trigger="none", continuation_risk="0", unmet_need="none")
            items.append(item)
        return {"items": items}, "fake-model"


def write_raw():
    rows = []
    for i, (src, text) in enumerate(SYN * 3):
        rows.append({"id": f"T-{i}", "source": src, "kind": "post", "title": "", "text": f"{text} #{i}",
                     "url": f"https://example.com/{i}", "thread_id": str(i), "author_hash": "",
                     "rating": None, "score": random.randint(0, 9), "published_at": "unknown",
                     "captured_at": "2026-10-25", "query": ""})
    with (C.RAW / "synthetic.jsonl").open("w") as f:
        for r in rows:
            f.write(json.dumps(r) + "\n")


class Pipeline(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        from engine import analyze, prefilter
        write_raw()
        cls.corpus = prefilter.run()
        CL.classify(cls.corpus, client=FakeLLM(), log=lambda *a: None)
        cls.analysis = analyze.run()
        from engine import golden
        cls.sheet = golden.make_sheet(8)

    def test_prefilter_drops_non_sip(self):
        self.assertTrue(all("US stocks" not in r["text"] for r in self.corpus))
        self.assertEqual(len(self.corpus), 21)

    def test_all_coded_and_passed(self):
        done = CL.load_done()
        self.assertEqual(len(done), 21)
        self.assertTrue(all(d["status"] == "passed_first_try" for d in done.values()))

    def test_opportunities_ranked(self):
        opps = self.analysis["opportunities"]
        self.assertTrue(opps)
        self.assertEqual(opps[0]["score_index"], 100)
        self.assertIn("reliable_payments", {o["need"] for o in opps})

    def test_gate_rejects_invented_quote(self):
        rec = {"id": "X", "title": "", "text": "My SIP failed"}
        item = FakeLLM(break_quote=True).chat_json("", json.dumps({"records": [
            {"id": "X", "text": "My SIP failed"}]}))[0]["items"][0]
        self.assertIn("quote is not an exact substring of the record", CL.validate(item, rec))

    def test_gate_rejects_bad_enum_and_inconsistency(self):
        rec = {"id": "X", "title": "", "text": "I stopped my SIP last month"}
        item = {f: ([next(iter(v["values"]))] if v["multi"] else next(iter(v["values"])))
                for f, v in CB.FIELDS.items()}
        item.update(relevance="core", journey_stage="stopped", continuation_risk="1",
                    quote="stopped my SIP", need_text="", reason="x")
        self.assertTrue(any("requires continuation_risk=3" in e for e in CL.validate(dict(item), rec)))
        item["journey_stage"] = "gave_up"
        self.assertTrue(any("journey_stage" in e for e in CL.validate(dict(item), rec)))

    def test_ask_declines_advice(self):
        res = ASK.ask("Which is the best fund for SIP?", [])
        self.assertEqual(res["status"], "declined")

    def test_ask_insufficient_evidence(self):
        res = ASK.ask("cryptocurrency staking rewards", [], client=FakeLLM())
        self.assertEqual(res["status"], "insufficient")

    def test_citation_check(self):
        ev = [{"id": "A", "title": "", "text": "The mandate failed silently"}]
        good, failed = ASK.check_citations({"claims": [
            {"text": "Mandates fail", "evidence": [{"id": "A", "quote": "mandate failed silently"}]},
            {"text": "Invented", "evidence": [{"id": "A", "quote": "this is not there"}]},
            {"text": "Most (87%) users", "evidence": [{"id": "A", "quote": "mandate failed silently"}]},
        ]}, ev)
        self.assertEqual(len(good), 1)
        self.assertEqual(len(failed), 2)
        good2, _ = ASK.check_citations({"claims": [
            {"text": "Mandates fail", "evidence": [{"id": "A", "quote": "Mandate failed, silently!"}]}]}, ev)
        self.assertEqual(len(good2), 1)  # punctuation and case differences are tolerated

    def test_golden_sheet_is_blind(self):
        from openpyxl import load_workbook

        from engine import golden
        ws = load_workbook(self.sheet)["label_here"]
        header = [c.value for c in ws[1]]
        self.assertNotIn("quote", header)
        self.assertTrue(all(ws.cell(2, header.index(f) + 1).value in (None, "") for f in golden.GOLD_FIELDS))

    def test_golden_score(self):
        import pandas as pd

        from engine import golden
        ids = json.loads((C.EVAL / "golden_ids.json").read_text())
        done = CL.load_done()
        rows = []
        for i in ids:
            row = {"id": i, **{f: done[i][f] for f in golden.GOLD_FIELDS}}
            rows.append(row)
        rows[0]["journey_stage"] = "stopped"  # one deliberate disagreement
        pd.DataFrame(rows).to_csv(C.EVAL / "golden_labels.csv", index=False)
        rep = golden.score()
        self.assertEqual(rep["items_labelled"], len(ids))
        self.assertEqual(rep["items_fully_agreed"], len(ids) - 1)


if __name__ == "__main__":
    unittest.main()
